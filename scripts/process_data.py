"""
Process multi-modal data for Phase 2 POC.
Extracts text from PDFs and Word docs, generates embeddings for images using CLIP.
"""

import json
import os
from pathlib import Path
from typing import List, Dict, Any
import argparse
from dotenv import load_dotenv

# Load environment variables
load_dotenv()
from tqdm import tqdm
import pdfplumber
from docx import Document
from PIL import Image
import torch
from transformers import CLIPProcessor, CLIPModel
from sentence_transformers import SentenceTransformer


class MultiModalProcessor:
    """Process multi-modal documents and generate embeddings."""

    def __init__(self):
        """Initialize processors and models."""
        print("Loading embedding models...")

        # Text embeddings (same as Phase 1)
        self.text_model = SentenceTransformer('all-MiniLM-L6-v2')
        print(f"✓ Loaded text model: all-MiniLM-L6-v2 (384 dims)")

        # Image embeddings using CLIP
        self.clip_model = CLIPModel.from_pretrained("openai/clip-vit-base-patch32")
        self.clip_processor = CLIPProcessor.from_pretrained("openai/clip-vit-base-patch32")
        self.device = "cuda" if torch.cuda.is_available() else "cpu"
        self.clip_model.to(self.device)
        print(f"✓ Loaded CLIP model: clip-vit-base-patch32 (512 dims)")
        print(f"  Device: {self.device}")

    def extract_text_from_pdf(self, pdf_path: str) -> Dict[str, Any]:
        """Extract text from PDF document."""
        try:
            with pdfplumber.open(pdf_path) as pdf:
                # Extract text from all pages
                text_content = []
                tables_content = []

                for page in pdf.pages:
                    # Extract text
                    page_text = page.extract_text()
                    if page_text:
                        text_content.append(page_text)

                    # Extract tables
                    tables = page.extract_tables()
                    if tables:
                        for table in tables:
                            # Convert table to text
                            table_text = "\n".join([" | ".join([str(cell) if cell else "" for cell in row]) for row in table])
                            tables_content.append(table_text)

                full_text = "\n\n".join(text_content)
                tables_text = "\n\n".join(tables_content) if tables_content else ""

                return {
                    'text': full_text,
                    'tables': tables_text,
                    'num_pages': len(pdf.pages),
                    'has_tables': len(tables_content) > 0,
                    'success': True
                }
        except Exception as e:
            return {
                'text': '',
                'tables': '',
                'num_pages': 0,
                'has_tables': False,
                'success': False,
                'error': str(e)
            }

    def extract_text_from_word(self, docx_path: str) -> Dict[str, Any]:
        """Extract text from Word document."""
        try:
            doc = Document(docx_path)

            # Extract paragraphs
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]

            # Extract tables
            tables_content = []
            for table in doc.tables:
                table_text = "\n".join([" | ".join([cell.text for cell in row.cells]) for row in table.rows])
                tables_content.append(table_text)

            full_text = "\n\n".join(paragraphs)
            tables_text = "\n\n".join(tables_content) if tables_content else ""

            return {
                'text': full_text,
                'tables': tables_text,
                'num_paragraphs': len(paragraphs),
                'num_tables': len(doc.tables),
                'success': True
            }
        except Exception as e:
            return {
                'text': '',
                'tables': '',
                'num_paragraphs': 0,
                'num_tables': 0,
                'success': False,
                'error': str(e)
            }

    def generate_text_embedding(self, text: str) -> List[float]:
        """Generate embedding for text using sentence-transformers."""
        if not text or not text.strip():
            # Return zero vector for empty text
            return [0.0] * 384

        embedding = self.text_model.encode(text, convert_to_tensor=False)
        return embedding.tolist()

    def generate_image_embedding(self, image_path: str) -> List[float]:
        """Generate embedding for image using CLIP."""
        try:
            image = Image.open(image_path).convert('RGB')

            # Process image
            inputs = self.clip_processor(images=image, return_tensors="pt")
            inputs = {k: v.to(self.device) for k, v in inputs.items()}

            # Generate embedding
            with torch.no_grad():
                image_features = self.clip_model.get_image_features(**inputs)

            # Normalize and convert to list
            embedding = image_features.cpu().numpy()[0]
            # Normalize
            embedding = embedding / (embedding**2).sum()**0.5

            return embedding.tolist()
        except Exception as e:
            print(f"Error processing image {image_path}: {str(e)}")
            # Return zero vector on error
            return [0.0] * 512

    def process_pdfs(self, pdfs_dir: Path, limit: int = None) -> List[Dict[str, Any]]:
        """Process all PDF documents."""
        print(f"\n{'='*60}")
        print("PROCESSING PDF DOCUMENTS")
        print(f"{'='*60}\n")

        pdf_files = list(pdfs_dir.glob("*.pdf"))
        if limit:
            pdf_files = pdf_files[:limit]

        processed_docs = []

        for pdf_path in tqdm(pdf_files, desc="Processing PDFs"):
            # Extract text
            extraction = self.extract_text_from_pdf(str(pdf_path))

            if not extraction['success']:
                print(f"✗ Failed to extract from {pdf_path.name}: {extraction.get('error', 'Unknown error')}")
                continue

            # Combine text and tables for embedding
            combined_text = extraction['text']
            if extraction['tables']:
                combined_text += "\n\n" + extraction['tables']

            # Generate embedding
            embedding = self.generate_text_embedding(combined_text[:2000])  # Limit to first 2000 chars

            # Create document record
            doc_record = {
                'id': pdf_path.stem,
                'filename': pdf_path.name,
                'type': 'pdf',
                'text': extraction['text'][:1000],  # Store first 1000 chars
                'full_text': combined_text,
                'num_pages': extraction['num_pages'],
                'has_tables': extraction['has_tables'],
                'embedding': embedding,
                'embedding_dim': len(embedding)
            }
            processed_docs.append(doc_record)

        print(f"\n✓ Processed {len(processed_docs)} PDF documents")
        return processed_docs

    def process_word_docs(self, word_dir: Path, limit: int = None) -> List[Dict[str, Any]]:
        """Process all Word documents."""
        print(f"\n{'='*60}")
        print("PROCESSING WORD DOCUMENTS")
        print(f"{'='*60}\n")

        word_files = list(word_dir.glob("*.docx"))
        if limit:
            word_files = word_files[:limit]

        processed_docs = []

        for docx_path in tqdm(word_files, desc="Processing Word docs"):
            # Extract text
            extraction = self.extract_text_from_word(str(docx_path))

            if not extraction['success']:
                print(f"✗ Failed to extract from {docx_path.name}: {extraction.get('error', 'Unknown error')}")
                continue

            # Combine text and tables for embedding
            combined_text = extraction['text']
            if extraction['tables']:
                combined_text += "\n\n" + extraction['tables']

            # Generate embedding
            embedding = self.generate_text_embedding(combined_text[:2000])  # Limit to first 2000 chars

            # Create document record
            doc_record = {
                'id': docx_path.stem,
                'filename': docx_path.name,
                'type': 'word',
                'text': extraction['text'][:1000],  # Store first 1000 chars
                'full_text': combined_text,
                'num_paragraphs': extraction['num_paragraphs'],
                'num_tables': extraction['num_tables'],
                'embedding': embedding,
                'embedding_dim': len(embedding)
            }
            processed_docs.append(doc_record)

        print(f"\n✓ Processed {len(processed_docs)} Word documents")
        return processed_docs

    def process_images(self, images_dir: Path, metadata_file: Path, limit: int = None) -> List[Dict[str, Any]]:
        """Process all images with CLIP."""
        print(f"\n{'='*60}")
        print("PROCESSING IMAGES WITH CLIP")
        print(f"{'='*60}\n")

        # Load metadata
        with open(metadata_file) as f:
            metadata = json.load(f)

        if limit:
            metadata = metadata[:limit]

        processed_images = []

        for meta in tqdm(metadata, desc="Processing images"):
            # Get base directory from environment variable
            base_dir = Path(os.getenv("DATA_OUTPUT_DIR", "./data/multimodal")).parent
            image_path = base_dir / meta['path']

            if not image_path.exists():
                print(f"✗ Image not found: {image_path}")
                continue

            # Generate CLIP embedding
            embedding = self.generate_image_embedding(str(image_path))

            # Also generate text embedding from description
            text_embedding = self.generate_text_embedding(meta['description'])

            # Create image record
            image_record = {
                'id': meta['image_id'],
                'filename': meta['filename'],
                'type': 'image',
                'damage_type': meta['damage_type'],
                'severity': meta['severity'],
                'claim_id': meta['claim_id'],
                'policy_id': meta['policy_id'],
                'location': meta['location'],
                'description': meta['description'],
                'image_embedding': embedding,
                'text_embedding': text_embedding,
                'image_embedding_dim': len(embedding),
                'text_embedding_dim': len(text_embedding)
            }
            processed_images.append(image_record)

        print(f"\n✓ Processed {len(processed_images)} images")
        return processed_images


def main():
    """Main entry point."""
    parser = argparse.ArgumentParser(description='Process multi-modal data')
    parser.add_argument('--type', choices=['pdf', 'word', 'images', 'all'],
                       default='all', help='Type of data to process')
    parser.add_argument('--limit', type=int, default=None,
                       help='Limit number of files to process (for testing)')

    args = parser.parse_args()

    # Change to scripts directory if not already there
    script_dir = Path(__file__).parent
    os.chdir(script_dir)

    # Initialize processor
    processor = MultiModalProcessor()

    # Define paths
    data_dir = Path(os.getenv("DATA_OUTPUT_DIR", "./data/multimodal"))
    pdfs_dir = data_dir / "pdfs"
    word_dir = data_dir / "word"
    images_dir = data_dir / "images"
    metadata_file = images_dir / "image_metadata.json"

    # Process data
    results = {}

    if args.type in ['pdf', 'all']:
        if pdfs_dir.exists():
            results['pdfs'] = processor.process_pdfs(pdfs_dir, args.limit)
        else:
            print(f"⚠️  PDF directory not found: {pdfs_dir}")

    if args.type in ['word', 'all']:
        if word_dir.exists():
            results['word_docs'] = processor.process_word_docs(word_dir, args.limit)
        else:
            print(f"⚠️  Word directory not found: {word_dir}")

    if args.type in ['images', 'all']:
        if images_dir.exists() and metadata_file.exists():
            results['images'] = processor.process_images(images_dir, metadata_file, args.limit)
        else:
            print(f"⚠️  Images directory or metadata not found")

    # Save processed data
    output_dir = data_dir / "processed"
    output_dir.mkdir(exist_ok=True)

    for data_type, data in results.items():
        output_file = output_dir / f"{data_type}_processed.json"
        with open(output_file, 'w') as f:
            json.dump(data, f, indent=2)
        print(f"\n✓ Saved {len(data)} {data_type} to {output_file}")

    # Print summary
    print(f"\n{'='*60}")
    print("PROCESSING SUMMARY")
    print(f"{'='*60}")
    for data_type, data in results.items():
        print(f"{data_type:15s}: {len(data):4d} documents processed")
    print(f"{'='*60}\n")


if __name__ == "__main__":
    main()
